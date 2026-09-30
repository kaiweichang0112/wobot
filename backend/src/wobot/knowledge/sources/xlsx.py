"""Read the product catalog workbook into rows of untouched cell text."""

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.worksheet.hyperlink import Hyperlink

SHEET_NAME = "產品整理"
HEADERS = (
    "產品名稱",
    "公司名稱",
    "公司地址",
    "連絡電話",
    "產品網址",
    "主要功能",
    "使用方式",
    "100字簡介",
    "產品第一層分類",
    "產品第二層分類",
    "本計畫導入年分",
)
URL_HEADER = "產品網址"


class CatalogSchemaError(Exception):
    """The workbook no longer has the sheet or columns that ingestion maps."""


@dataclass(frozen=True)
class CatalogRow:
    # The Excel row number, meaningful only within the snapshot it was read from.
    row_number: int
    # Header → cell text exactly as stored; cleaning belongs to normalization.
    cells: dict[str, str | None]
    # Target of the hyperlink object on the URL cell. The text alone can differ from it,
    # and one row in the catalog has text but no hyperlink object.
    product_url_target: str | None


def read_catalog(content: bytes) -> list[CatalogRow]:
    """Parse the workbook, failing on any change to the sheet name or the columns."""
    # Not read-only mode: it skips hyperlinks, which the URL column needs.
    workbook = load_workbook(BytesIO(content), data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise CatalogSchemaError(f"sheet {SHEET_NAME!r} not found in {workbook.sheetnames}")
    sheet = workbook[SHEET_NAME]

    headers = [_text(cell.value) for cell in sheet[1]]
    while headers and headers[-1] in (None, ""):
        headers.pop()
    if tuple(headers) != HEADERS:
        # Strict on purpose: a renamed, moved or added column needs a person to map it.
        raise CatalogSchemaError(f"expected columns {list(HEADERS)}, found {headers}")

    url_column = HEADERS.index(URL_HEADER)
    rows = []
    for cells in sheet.iter_rows(min_row=2, max_col=len(HEADERS)):
        values = [_text(cell.value) for cell in cells]
        if all(value is None or not value.strip() for value in values):
            continue  # formatting-only rows below the data
        rows.append(
            CatalogRow(
                row_number=cells[0].row,
                cells=dict(zip(HEADERS, values, strict=True)),
                product_url_target=_hyperlink_url(cells[url_column].hyperlink),
            )
        )
    return rows


def _hyperlink_url(hyperlink: Hyperlink | None) -> str | None:
    """The external URL of a hyperlink, with its fragment put back.

    Excel stores `https://host/page#anchor` as target `https://host/page` plus location
    `anchor`; a link with a location only points inside the workbook and is no URL.
    """
    if hyperlink is None or not hyperlink.target:
        return None
    if hyperlink.location:
        return f"{hyperlink.target}#{hyperlink.location}"
    return hyperlink.target


def _text(value: object) -> str | None:
    """Cell value as text; Excel may store years or phone numbers as numbers."""
    match value:
        case None | str():
            return value
        case bool():
            return "TRUE" if value else "FALSE"
        case float() if value.is_integer():
            return str(int(value))
        case datetime() | date():
            return value.isoformat()
        case _:
            return str(value)
